from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "reports" / "Algoritmik_Finans_Arastirma_Platformu_TUBITAK_BIST_Teknik_Rapor_TR.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)
FIG = Path(__file__).resolve().parent / "figures"
FIG.mkdir(parents=True, exist_ok=True)
SHOT_ROOT = ROOT / "IndicatorCodeExamples" / "OptimalSettingsForLiqNarrative" / "Muhteşem4lü" / "RealDeals"

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.66)
sec.left_margin = Inches(0.85)
sec.right_margin = Inches(0.85)
sec.header_distance = Inches(0.32)
sec.footer_distance = Inches(0.32)
sec.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(25, 25, 25)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size, before, after in [("Title", 26, 0, 12), ("Subtitle", 13, 0, 9), ("Heading 1", 16, 15, 8), ("Heading 2", 11.5, 10, 5)]:
    s = styles[name]
    s.font.name = "Arial"
    s.font.size = Pt(size)
    s.font.bold = name != "Subtitle"
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True
if "Caption" not in styles:
    styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
styles["Caption"].font.name = "Arial"
styles["Caption"].font.size = Pt(9)
styles["Caption"].font.bold = True
styles["Caption"].font.color.rgb = RGBColor(30, 30, 30)
styles["Caption"].paragraph_format.space_before = Pt(8)
styles["Caption"].paragraph_format.space_after = Pt(5)
styles["Caption"].paragraph_format.keep_with_next = True


def add(text, style=None):
    para = doc.add_paragraph(style=style)
    para.add_run(text)
    para.paragraph_format.widow_control = True
    return para


def heading(text, level=1):
    doc.add_heading(text, level)


def bullets(items):
    for item in items:
        add(item, "List Bullet")


def caption(text):
    add(text, "Caption")


def set_cell(cell, value, bold=False):
    cell.text = str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for para in cell.paragraphs:
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.line_spacing = 1.04
        for run in para.runs:
            run.font.name = "Arial"
            run.font.size = Pt(8.6)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(255, 255, 255) if bold else RGBColor(25, 25, 25)
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:w"), "90")
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcpr.append(mar)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths:
        for row in t.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    for i, value in enumerate(headers):
        set_cell(t.rows[0].cells[i], value, True)
    for r, values in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(values):
            set_cell(cells[i], value)
    for ri, row in enumerate(t.rows):
        for cell in row.cells:
            tcpr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "213B58" if ri == 0 else ("F1F5F8" if ri % 2 == 0 else "FFFFFF"))
            tcpr.append(shd)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                node = OxmlElement("w:" + edge)
                node.set(qn("w:val"), "single")
                node.set(qn("w:sz"), "4")
                node.set(qn("w:color"), "D9D9D9")
                borders.append(node)
            tcpr.append(borders)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def page_break():
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def picture(path, title, note, width=6.45):
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    caption(title)
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_together = True
    para.add_run().add_picture(str(path), width=Inches(width))
    note_para = add(note)
    note_para.paragraph_format.space_after = Pt(9)
    for run in note_para.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(75, 75, 75)


def architecture_figure():
    labels = ["Lisanslı veri\nve referans", "Normalize olay\nkalite ve UTC", "Strateji\naday sinyal", "Rejim ML\nve risk kapısı", "Audit veya\ntest icrası", "Olay bazlı\nmutabakat"]
    fig, ax = plt.subplots(figsize=(11.0, 2.35), dpi=175)
    fig.patch.set_facecolor("white")
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 2)
    ax.axis("off")
    colors = ["#e8f0f6", "#dceaf4", "#d5e6ef", "#bed8e7", "#a9c9dc", "#8fb8d0"]
    for i, (label, color) in enumerate(zip(labels, colors)):
        x = 0.12 + i * 2.0
        ax.add_patch(FancyBboxPatch((x, 0.53), 1.72, 0.94, boxstyle="round,pad=0.04,rounding_size=0.09", linewidth=1, edgecolor="#54728a", facecolor=color))
        ax.text(x + 0.86, 1, label, ha="center", va="center", fontsize=9, color="#172a38", fontweight="bold")
        if i < 5:
            ax.add_patch(FancyArrowPatch((x + 1.74, 1), (x + 1.96, 1), arrowstyle="-|>", mutation_scale=15, lw=1.2, color="#54728a"))
    path = FIG / "architecture.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def research_figure():
    labels = ["Ayarları dondur", "Geliştirme\nve maliyet", "Kronolojik\ndoğrulama", "Kilitli\nson dönem", "Gölge izleme\nve olay kaydı"]
    fig, ax = plt.subplots(figsize=(10, 2.5), dpi=175)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 2)
    ax.axis("off")
    for i, label in enumerate(labels):
        x = 0.08 + 2 * i
        ax.add_patch(FancyBboxPatch((x, .58), 1.68, .84, boxstyle="round,pad=0.04,rounding_size=0.07", facecolor="#f2f5f7" if i % 2 == 0 else "#e0ebf2", edgecolor="#668298"))
        ax.text(x + .84, 1, label, ha="center", va="center", fontsize=9, color="#1f3443")
        if i < 4:
            ax.add_patch(FancyArrowPatch((x + 1.72, 1), (x + 1.95, 1), arrowstyle="-|>", mutation_scale=15, color="#668298"))
    path = FIG / "research.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def comparison_figure():
    labels = ["USDJPY\nLojistik", "USDJPY\nHGB", "JP225\nLojistik", "JP225\nHGB"]
    values = [-0.39, -1.39, -1.01, -1.75]
    fig, ax = plt.subplots(figsize=(8.8, 3.1), dpi=175)
    bars = ax.bar(labels, values, color="#758da1", width=.54)
    ax.axhline(0, color="#293745", lw=1)
    ax.set_ylabel("Medyan fark  R")
    ax.set_ylim(-2.1, .35)
    ax.spines[["top", "right"]].set_visible(False)
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width()/2, value-.08, f"{value:.2f}R", ha="center", va="top", fontsize=9)
    fig.tight_layout()
    path = FIG / "r7_comparison.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


header = sec.header.paragraphs[0]
header.text = "Algoritmik Finans Araştırma Platformu | Teknik Ar-Ge Raporu"
header.style = normal
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(90, 90, 90)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("Cem Berk Çıracı  |  23 Eylül 2026  |  ")
run.font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

title = doc.add_paragraph(style="Title")
title.add_run("Algoritmik Finans Araştırma Platformu")
add("Teknik mimari, doğrulama yöntemi ve üniversite sanayi iş birliği önerisi", "Subtitle")
add("TÜBİTAK 1505 ön görüşmeleri ve Borsa İstanbul potansiyel paydaşlığı için teknik rapor")
add("Hazırlayan  Cem Berk Çıracı\nBelge kodu  ALGOFIN-AR-GE-TR-001\nSürüm  1.0\nTarih  23 Eylül 2026")
add("Bu belge başvuru kararı, ortaklık taahhüdü, yatırım tavsiyesi veya Borsa İstanbul üzerinde doğrulanmış getiri beyanı değildir. Mevcut yazılım kanıtlarını önerilen yeni araştırma ve pilot aşamalarından ayırır.")
page_break()

heading("İçindekiler")
for line in [
    "1 Yönetici özeti ve proje kararı", "2 Mevcut araştırma ortamı ve kanıt sınırları",
    "3 Borsa İstanbul bağlamında araştırma problemi", "4 Uçtan uca hedef mimari",
    "5 Veri sözleşmesi ve nedensellik", "6 Audit ve işlem simülasyonu",
    "7 Strateji araştırması, rejim ve yapay zekâ", "8 Emir yaşam döngüsü ve risk güvenliği",
    "9 Deney protokolü ve başarı ölçütleri", "10 İş paketleri ve ortakların rolleri",
    "11 Veri hakkı, fikri mülkiyet ve yayına hazırlık", "12 Başvuru yolu ve karar kapıları",
    "Ek A Kaynak kod kanıt dizini", "Kaynakça",
]:
    add(line)
page_break()

heading("1 Yönetici özeti ve proje kararı")
add("Bu önerinin konusu, yalnızca kârlı sinyal üreten bir bot değildir. Yaklaşık bir yıllık geliştirme ile oluşmuş MT5 tabanlı algoritmik araştırma ortamının; çoklu strateji çekirdeklerini, sürümlenmiş ayar kütüphanelerini, tarihsel audit motorlarını, makine öğrenmesi filtre deneylerini, CSV tabanlı yayın katmanını ve farklı işlem platformlarındaki dinleyicileri aynı deney zincirinde buluşturmasıdır. Araştırma değeri, bir stratejinin neden ve hangi veri kalitesi koşullarında başarılı göründüğünü yeniden üretilebilir biçimde sınayabilmesidir [K01–K11].")
add("Mevcut prototipte Liquidity Narrative, opening-range breakout (ORB), NAS SATS ve hareketli ortalama kesişimi (MA) aileleri vardır. Risk geometrisi, rejim sınıflaması, sinyal ve emir ayrımı, varlığa özgü seans/oynaklık parametreleri, tarihsel tarama ve bazı uçtan uca listener akışları kod düzeyinde mevcuttur. Ancak bu bileşenlerin bulunması, hepsinin her brokerda eşdeğer veya Borsa İstanbul’da üretim hazır olduğu anlamına gelmez. Raporun önerdiği Ar-Ge, mevcut altyapıyı lisanslı BIST verisi ve kurumca sağlanacak test koşulları üzerinde bağımsız biçimde sınamaktır.")
add("Önerilen 18 aylık ortak çalışma, üç çıktıya odaklanır: (i) veri kaynaklarını ve sinyal kullanılabilirlik zamanını açıkça tanımlayan araştırma veri sözleşmesi; (ii) bar/tick belirsizliğini gizlemeyen, canlı emir kayıtlarıyla olay bazında mutabakat yapabilen audit; (iii) ileriye dönük, maliyet sonrası ve çoklu deneme yanlılığı gözetilmiş strateji/filtre değerlendirme protokolü. Dördüncü çıktı, yalnızca kurum izni ve teknik sertifikasyon süreci uygun olursa, BISTECH test ortamında güvenlik odaklı entegrasyon gösterimidir [1–4].")
caption("Tablo 1  Mevcut varlık ile araştırmada hedeflenen kanıt düzeyi")
table(["Alan", "Bugün gözlenen", "Projenin kanıtlayacağı"], [
    ["Strateji", "Dört EA ailesi, kaydedilmiş ayarlar ve rejim/filtre kombinasyonları", "BIST ürünlerinde maliyet sonrası ileri dönem geçerliliği veya açık biçimde başarısızlık"],
    ["Audit", "Hızlı tarihsel tekrar oynatma, bazı ailelerde tick veya M1/bar belirsizlik işaretleri", "BIST veri kapsaması, etiket doğruluğu, olay düzeyinde canlı/test mutabakatı"],
    ["İşlem hattı", "MT5 yayıncıları ve cTrader/NinjaTrader/Quantower dinleyicileri", "Kanal sahipliği, yeniden başlama, tekrar, gecikme ve emir durum makinesi testleri"],
    ["Yapay zekâ", "Strateji bazlı filtre, kronolojik ayrım ve reddedilen model kaydı", "Sabit deney protokolü, daha güçlü baz çizgileri, ileri dönem kalibrasyon ve seçim yanlılığı ölçümü"],
], [1.1, 2.55, 2.95])
add("Başvuru kararı: TÜBİTAK 1505 uygun bir ön çerçevedir; program, Türkiye’de yerleşik müşteri kuruluş ile üniversite/araştırma yürütücüsünün ortak başvurusunu ve başvurudan önce iş birliği sözleşmesini öngörür [1]. Borsa İstanbul burada doğrulanmış taraf değil, olası veri, piyasa uzmanlığı ve test ortamı paydaşıdır. Müşteri kuruluş rolü ve veri hakları görüşmede netleşmeden resmî ortaklık beyanı yapılmamalıdır.")
picture(SHOT_ROOT / "SolAdvanced.png", "Şekil 1  Sol Advanced prototipinde rejim ve audit paneli", "Arşiv ekran görüntüsü  24 Temmuz 2026. Gösterilen istatistikler tarihsel prototip çıktısıdır; BIST verisi veya bağımsız ileri test sonucu değildir.", 5.85)

heading("2 Mevcut araştırma ortamı ve kanıt sınırları")
heading("2.1 Çekirdek stratejiler ve konfigürasyon", 2)
add("Liquidity Narrative ailesi ERL/IRL, sweep, FVG ve rejim temelli kural katmanlarını; ORB ailesi açılış aralığı, seans ve risk geometrisini; NAS ailesi bağımsız sinyal/filtre ve çoklu işlem mantığını; MA ailesi çeşitli kesişim ve trend modlarını içerir [K01–K04]. Bunlar tek bir evrensel alpha modeli değil, aynı araştırma disiplinine bağlanmış farklı hipotez aileleridir. Kullanıcı ayarlarının strateji kimliğiyle saklanması, aynı parametrelerin tekrar çağrılmasına yardımcı olur; modelin tarihsel araştırma sürecinde kaç kez seçildiğini tek başına kayıt altına almaz.")
picture(SHOT_ROOT / "GiveATry" / "UseThis" / "ORB0.5" / "Nikkei" / "NikkeiHighAccuracy" / "JP2251RSetup.png", "Şekil 2  ORB tarihsel audit ve strateji parametre ekranı", "Arşiv ekran görüntüsü  20 Ağustos 2026. 1Y sinyal, W/L, net R, rejim ve maliyet alanlarını gösterir. Seçilmiş geçmiş performans yeni piyasada genellenemez.", 5.0)
heading("2.2 Hızlı tarama ile sıkı doğrulama ayrımı", 2)
add("Dashboard audit motorları çok sayıda yapılandırmayı hızlı taramaya yarar. Repo içindeki doğrulama yaklaşımı, bu sonuçların örnek olaylarını MT5 Strategy Tester işlem kayıtlarıyla karşılaştırmayı da tarif eder [K09]. Buna rağmen tüm dönemlerin tick düzeyinde, aynı emir defteri ve aynı komisyon/sürçme koşullarında simüle edildiği söylenemez. MA aile dokümanı M1 Bid OHLC ve kaydedilmiş spread kullandığını, aynı mumdaki SL/TP sırası belirsizse DATA döndürdüğünü açıkça yazar [K08]. NAS strateji kütüphanesi araştırma notu da bazı HCBE yapılandırmalarının sıkı CSV eğitim etiketine alınmadığını belirtir [K10].")
picture(SHOT_ROOT / "GiveATry" / "UseThis" / "ORB0.5" / "Crypto" / "Nas" / "BTCNAS.png", "Şekil 3  NAS prototipinde audit ve hesap riski kontrolleri", "Arşiv ekran görüntüsü  28 Ağustos 2026. BTC bağlamındaki bu örnek BIST işlem sonucu değildir; platformun kontrol ve raporlama yüzeyini gösterir.", 5.8)
heading("2.3 Makine öğrenmesi ve başarısız deneylerin değeri", 2)
add("ORB ML protokolünde satır birimi kapanmış işlemdir; etiket olan çıkış zamanı/net R gelecekten özelliklere sızdırılmaz, eğitim ve test kronolojik ayrılır [K11]. Daha sonraki R7 ve R8 araştırmaları otomatik üstünlük göstermemiştir: 49 kayıtlı stratejinin eşlenik karşılaştırmasında R7’nin tüm ailelerde medyan farkı negatif kalmıştır; R8 çoklu test düzeltmesinden sonra anlamlı üstün aday raporlamamıştır [K12–K13]. Bu negatif sonuçlar, önerinin bilimsel gücünü artırır: platform yalnızca iyi görünen grafikleri değil, reddedilen hipotezleri de arşivleyebilir. R7/R8 raporlarındaki seçilmiş pozitif örnekler, bağımsız BIST performansı değildir.")
heading("2.4 Çoklu platform sınırı", 2)
add("Kaynak ağacında cTrader, NinjaTrader ve Quantower dinleyicileri bulunur [K05–K07]. CSV köprüsü ürün ve platform bağlamında pratik bir prototiptir; farklı enstrümanların kontrat büyüklüğü, fiyat kotasyonu ve işlem saatleri birbirinden farklıdır. JP225 CFD’de iyi görünen bir stop/target veya risk yüzdesi BIST payı, VİOP kontratı ya da başka platformda otomatik eşdeğer sayılamaz. Yeni projede sembol, tick değeri, lot/kontrat ve oturum bilgisini referans veriyle doğrulayan ayrı bir adaptör gereklidir.")

heading("3 Borsa İstanbul bağlamında araştırma problemi")
add("Borsa İstanbul verisi lisanslı veri dağıtım kuruluşları, farklı yayın paketleri ve tarihsel DataStore kanalları üzerinden sunulur [2, 3]. Mevcut broker CFD akışlarını BIST işlemi olarak yeniden adlandırmak yöntemsel olarak yanlıştır: referans fiyat, işlem takvimi, likidite, mikro yapı ve kurumca uygulanan işlem maliyetleri farklıdır. Araştırma, veri ürününün kapsamı ve kullanım izni netleşmeden başlamamalıdır.")
add("Birinci araştırma sorusu, aynı strateji kuralının farklı veri çözünürlüklerinde ve belirsiz işlem sıralarında ne kadar kararlı olduğudur. İkincisi, rejim ve yapay zekâ filtresinin belirli BIST ürünlerinde maliyet sonrası net fayda sağlayıp sağlamadığıdır. Üçüncüsü, prototip CSV/listener hattından üretim benzeri emir olaylarına geçerken tekrar, gecikme, kopma veya yeniden başlatmanın risk sınırlarını aşıp aşmadığıdır. Dördüncüsü, varsayılan piyasa rejimlerinin başka döneme ve ürüne taşınabilir olup olmadığıdır.")
caption("Tablo 2  BIST pilotu için veri ve izin kapıları")
table(["Girdi", "Önerilen kullanımı", "Başlamadan doğrulanacak"], [
    ["Fiyat ve işlem", "Seans içi sinyal ve maliyet çalışması", "Lisans, geçmiş kapsama, düzeltmeler, gecikme, yeniden dağıtım hakkı"],
    ["Emir defteri / olay", "Varsa mikro yapı ve gerçekleşebilirlik analizi", "Seviye derinliği, zaman damgası, üyelik/anonimlik, erişim bedeli"],
    ["Referans veri", "Ürün kodu, tick, lot, takvim ve sözleşme vadesi", "Kurumsal işlem ve vadeli kontrat devamlılık kuralları"],
    ["Test ortamı", "Emir durum makinesi ve risk enjeksiyonu", "Kullanıcı hesabı, yetki, protokol, sertifikasyon sınırı"],
], [1.25, 2.25, 3.1])
add("BISTECH test ortamlarının erişim yolu ve kullanıcı tahsisi Borsa İstanbul tarafından açıklanır [4]. FIXAPI/OUCH canlı emir iletimi için üye ve yazılım sertifikasyon koşulları ayrıca belirtilir [5]. Bu rapor, doğrudan canlı BIST emri gönderen hazır entegrasyon bulunduğunu iddia etmez.")

heading("4 Uçtan uca hedef mimari")
picture(architecture_figure(), "Şekil 4  Önerilen araştırma veri ve karar zinciri", "Temsili mimari çizimi. BIST veri adaptörü ve BISTECH test entegrasyonu mevcut çalışan bileşen olarak değil, önerilen iş paketi olarak gösterilmiştir.")
caption("Tablo 3  Önerilen zincirin veri ve karar sözleşmeleri")
table(["Aşama", "Sözleşme", "Çıktı"], [
    ["1 Veri kabul", "Lisanslı kaynak, referans veri, UTC ve kaynak zamanı", "Değişmez ham veri + kalite bayrağı"],
    ["2 Normalizasyon", "Ürün, seans, tick, maliyet, kurumsal işlem", "Sürüm numaralı olay akışı"],
    ["3 Sinyal", "Kapanmış/verisi gelmiş bar; strateji ve ayar kimliği", "Nedensel aday olay"],
    ["4 Politika", "Rejim, ML, risk ve izin kapıları", "İzin/red nedeni ve model kimliği"],
    ["5 Simülasyon", "Emir, kısmi dolum, iptal, koruma ve maliyet", "Olay düzeyinde audit kaydı"],
    ["6 Mutabakat", "Test/canlı emir ve işlem akışı", "Sapma, incident ve kill-switch kararı"],
], [1.2, 3.3, 2.1])
add("Her karar, event_id, strategy_id, config_hash, model_hash, source_symbol, venue_symbol, source_time, available_time, decision_time, data_version ve reason_code ile ilişkilendirilir. Böylece “sinyal ne zaman bilinebilirdi?” sorusu sonradan yanıtlanabilir. Ölçülmüş bir bid/ask yoksa spread varsayımı açıkça etiketlenir; kaynak boşluğunda tahmini tick yaratılmaz.")
heading("4.1 Ayrı kontrol ve veri düzlemleri", 2)
add("Araştırma düzleminde strateji, filtre ve audit tekrar oynatılır. İcra düzleminde emir gönderme, pozisyon sahipliği, koruma emirleri, iptal ve hesap risk sınırları bulunur. Filtre ON/OFF gibi deneysel durumlar, emir motorunun işlemi yanlış kanaldan sahiplenmesine izin vermemelidir. Gözlem ve deney kapıları, üretim işlemini otomatik etkinleştirmez.")
heading("4.2 Sürüm ve yeniden üretilebilirlik", 2)
add("Her raporlanan sonuç için kaynak commit/hash, derlenmiş EA sürümü, veri kümesi özeti, broker veya veri satıcısı, takvim sürümü, ücret modeli, ayar manifesti, tüm denenen aday sayısı ve reddedilen koşullar saklanır. Model yeniden eğitilince eski ağırlıkların üstüne yazılmaz. Strateji görselinden çıkarılmış ayar, ekran görüntüsü kanıtı olsa bile tarihsel veri ve kod sürümü olmadan bağımsız sonuç olarak sunulmaz.")

heading("5 Veri sözleşmesi ve nedensellik")
add("Araştırma tablosunun temel birimi, karar anında mevcut olan bir aday sinyaldir. Sinyal üretim anı ile işlemin başarı/başarısızlık etiketini belirleyen çıkış anı farklı alanlardır. Özellikler, available_time sonrasında oluşmuş yüksek/düşük, ATR, rejim veya daha sonra kesinleşmiş pivot bilgisini kullanamaz. HTF barı kapanmadan onun nihai değeri alt TF sinyaline taşınamaz.")
caption("Tablo 4  Asgari olay şeması ve kalite semantiği")
table(["Alan grubu", "Asgari içerik", "Kural"], [
    ["Kimlik", "event_id, strategy_id, channel, data_version", "Tekrarlar idempotent işlenir"],
    ["Zaman", "exchange_time, source_time, available_time, ingest_time", "UTC ve DST dönüşümü sürümlenir"],
    ["Piyasa", "bid, ask veya işlem fiyatı; hacim ve derinlik kalitesi", "Gözlenen ile tahmin ayrılır"],
    ["Karar", "direction, risk geometry, regime, veto reason", "Karar öncesi değerler dondurulur"],
    ["Sonuç", "fill, fee, exit reason, realized R, label quality", "Açık/ambiguous olay W/L yapılmaz"],
], [1.1, 3.0, 2.5])
add("BIST payları için bölünme, temettü, hak kullanımı ve işlem durdurma; VİOP için vade değişimi, tick/kontrat çarpanı ve fiyat limitleri veri katmanının parçasıdır. Bu alanlar veri ürününün gerçekten sunduğu kapsama göre kesinleştirilir. Veri kalitesi kesintileri ayrı deney senaryosu olarak raporlanır; eksik dönem performans tablosundan sessizce çıkarılmaz.")

heading("6 Audit ve işlem simülasyonu")
add("Audit motoru yalnızca W/L oranı hesaplamamalı; hangi fiyatın bilindiğini, işlemin hangi emirle açıldığını, koruma emrinin ne zaman etkinleştiğini ve hangi olayla kapandığını açıklamalıdır. Aynı zaman dilimindeki hem TP hem SL teması, daha ince veri yoksa belirsizdir. Sonucu kazanan ya da kaybeden olarak atamak yerine DATA / ambiguous intrabar olarak korumak araştırma dürüstlüğüdür. Exact tick de emir kuyruğu, piyasa etkisi ve broker reddi olmadan tam gerçekleşme garantisi vermez.")
add("Önerilen iki kademeli yöntem: geniş strateji uzayını mevcut hızlı motorla taramak; seçilen tüm adayları dondurulmuş parametrelerle daha sıkı tick/işlem maliyeti ve BIST test ortamı mutabakatına almak. Hızlı tarama sonuçları yayınlanabilir edge kanıtı değil, hipotez eleme aracıdır. Kademeler arasında tekil event_id ve risk geometri karşılaştırması yapılır.")
caption("Tablo 5  Audit doğrulama senaryoları")
table(["Deney", "Beklenen davranış", "Hata göstergesi"], [
    ["Aynı bar SL ve TP", "Zaman sırası yoksa DATA", "İyimser otomatik TP"],
    ["Ters sinyal çıkışı", "O anki uygulanabilir fiyatla kısmi R", "Tam TP veya tam SL yazılması"],
    ["Geç sinyal / yeniden başlama", "Süre ve sahiplik politikasıyla kontrol", "Eski sinyalden yeni pozisyon"],
    ["Maliyet şoku", "Net R ve retention duyarlılığı", "Sıfır ücretteki W/L ile yetinme"],
    ["Veri boşluğu", "Kalite bayrağı ve kapsam kaybı", "Sessiz doldurma veya geleceğe bakış"],
], [1.5, 2.6, 2.5])
add("Ücret modeli, aracılık komisyonu, borsa/takas bedeli, alış-satış farkı, gecikme ve pozisyon büyüklüğüne göre olası piyasa etkisini ayrı parametreler olarak taşır. Bilinmeyen bileşenler tahmini aralıkla duyarlılık analizine girer. Brüt ve net R, para birimi getirisi, işlem sayısı, minimum kümülatif PnL, kapanmış işlem zirve-dip drawdown ve gün içi açık zarar ayrı raporlanır; bu metrikler birbirinin yerine geçmez.")

heading("7 Strateji araştırması, rejim ve yapay zekâ")
heading("7.1 Strateji aileleri ortak laboratuvarda", 2)
add("Liq, ORB, NAS ve MA aynı test evreninde farklı hipotezler sağlar: likidite geri dönüşü, seans açılış kırılımı, trend/oynaklık bağımlı giriş ve hareketli ortalama durum geçişi. “En iyi strateji” tek bir tarihsel lider tablodan seçilmez; ürün, veri kaynağı, yıl, seans, maliyet varsayımı ve işlem sıklığına koşullu dağılım incelenir. Strateji geliştirme sırasında görülen veriler bağımsız ileri test diye etiketlenmez.")
heading("7.2 Rejim ölçümü", 2)
add("Rejimler yalnızca geriye dönük kazananları etiketlemek için değil, karar anında ölçülebilen piyasa özellikleriyle tanımlanacaktır. Trend yönü ve derinliği, gerçekleşen oynaklık, seans likiditesi, açılış boşluğu, hacim/dengesizlik ve H1/H2 gibi çoklu ufukların durumu aday özelliklerdir. Her sınıfın geçiş gecikmesi, sınıf büyüklüğü ve sınır yakınındaki kararsızlık raporlanır. Gelecekte gözlenen “bu dönem iyiymiş” etiketi canlı filtresine dönüştürülmez.")
heading("7.3 ML filtre ilkesi", 2)
add("Model yeni işlem icat etmez; mevcut setup için izin, veto veya belirsizlik kararı üretir. Karşılaştırma sırası filtresiz strateji, basit rejim kuralı, düzenlileştirilmiş doğrusal model, sınırlı ağaç ve ancak veri gerekçesi varsa sinir ağıdır. Başarı tek başına WR artışı değildir: net R, DD, karar sayısı, maliyet şoku, kalibrasyon, rejimler arası tutarlılık ve güven aralığı birlikte değerlendirilir. Eğitim verisinin azaldığı nadir rejimlerde modelin abstain etmesi kabul edilir.")
add("Model seçimi sırasında çok sayıda strateji, eşik, yıl ve rejim taranması aynı geçmişe tekrar bakma riskini artırır. Cawley ve Talbot model seçim kriterinin de aşırı uyabileceğini gösterir [6]; Bailey ve López de Prado’nun DSR yaklaşımı çoklu deneme ve normal olmayan getiri dağılımını dikkate alır [7]. Bu yöntemler mevcut sistemde bütünüyle uygulanmış gibi sunulmaz. Projede deneme defteri, kronolojik ileri bölme, örtüşen pozisyonların purge edilmesi, embargo, dondurulmuş final dönem ve bağımsız kâğıt işlem evresi birlikte kullanılacaktır.")
picture(comparison_figure(), "Şekil 5  R7 adaylarının mevcut modele karşı medyan net R farkı", "Kaynak: docs/ORB_Current_vs_R7_Comparison.md. Dört aday ailesi için 49 kayıtlı stratejide eşlenik tarihsel karşılaştırma; negatif medyan, tüm aileleri üretime alma iddiasını desteklemez. BIST testi değildir.")

heading("8 Emir yaşam döngüsü ve risk güvenliği")
add("Aynı stratejinin iki kanalı veya iki hesabı, aynı sembole gelse bile birbirinin pozisyonunu açıp kapatamaz. Bunun için pozisyon sahipliği yalnızca label metnine değil; hesap, venue, ürün, kanal, strateji ve emir kimliği birleşimine bağlanır. Dosya yayını atomik yazım, scoped latest seçimi, sequence ve acknowledgement testleriyle doğrulanır. Dosyanın okunması emrin kabul edildiği, dolduğu veya korunduğu anlamına gelmez.")
add("Risk motoru emir öncesinde sözleşme çarpanı, tick değeri, kur dönüşümü, pozisyon limiti, stop mesafesi, açık maruziyet ve hesap drawdown eşiğini kontrol eder. Emir sonrası gerçek fill ile yeniden hesaplar; aşım veya koruma emri reddinde fail-closed politika uygular. Bağlantı kaybı, yeniden başlatma, kısmi dolum, slippage, cancel/replace yarışı ve aynı sinyalin tekrar yayını ayrı enjeksiyon testleridir. Açık pozisyonun hafta sonu veya terminal yeniden başlamasında yeniden sahiplenilmesi ayrıca sınanır.")
add("BISTECH PTRM, borsanın işlem öncesi risk kontrollerine ilişkin ayrı bir katmandır [8]. Araştırma prototipindeki hesap koruması PTRM yerine geçmez. Canlı BIST emir aşaması ancak yetkili üye, veri/bağlantı koşulları, sertifikasyon ve hukuk/uyum onayları tamamlandıktan sonra düşünülebilir [4, 5].")

heading("9 Deney protokolü ve başarı ölçütleri")
picture(research_figure(), "Şekil 6  Ön kayıtlı araştırma ve doğrulama akışı", "Temsili yöntem şeması. Kilitli son dönem görüldükten sonra yapılan her değişiklik yeni deney kimliği alır.")
add("Hipotezler ve durdurma kuralları BIST verisini açmadan önce kaydedilir. Önce veri sözleşmesi ve baz çizgiler dondurulur. Sonra geliştirme döneminde parametre seçimi yapılır; doğrulama dönemi eşik seçimi için bir kez kullanılır; son dönem nihai rapor için kilitlenir. Son test görüldükten sonra yapılan değişiklik yeni deney numarası alır. Stratejiler aynı piyasa olaylarını paylaştığı için 50 strateji, 50 bağımsız deney sayılmaz. Gün/hafta bloklu yeniden örnekleme ve ürün bazında eşlenik farklar raporlanır.")
caption("Tablo 6  Önerilen pilot kabul ölçütleri")
table(["Ölçüt", "Hedef niteliği", "Kanıt yöntemi"], [
    ["Nedensel veri", "Kritik özelliklerde geleceğe erişim bulunmaması", "Zaman damgası ve sentetik gecikme testleri"],
    ["Replay eşdeğerliği", "Aynı veri/sürümle aynı karar dizisi", "Hash ve olay bazlı karşılaştırma"],
    ["Audit dürüstlüğü", "Ambiguous işlemin W/L’ye dönmemesi", "Çakışan TP/SL ve veri boşluğu testleri"],
    ["Kanal yalıtımı", "Enjekte edilen çapraz sahiplik girişimlerinin reddi", "İki hesap/iki kanal state-machine testi"],
    ["ML faydası", "İleri dönemde maliyet sonrası baz çizgiye eşlenik fark", "Güven aralığı, retention ve deneme sayısı"],
    ["BIST uyumu", "İzinli veride ürün/takvim/lot/tick tutarlılığı", "Referans veri ve kurum test ortamı mutabakatı"],
], [1.35, 2.3, 2.95])
add("Bu ölçütler bir getiri yüzdesi taahhüdü değildir. Son satırda başarısız sonuç, veri/strateji taşınabilirliğinin sınırını gösteren değerli bir bilimsel çıktı olabilir. Projenin asgari kabulü, doğru sonuç ve belirsizliğin ölçülmesidir; kâr sonucunun olumlu olması ön koşul değildir.")

heading("10 İş paketleri ve ortakların rolleri")
add("Aşağıdaki 18 aylık plan öneridir. Bütçe, uygun kurum ve veri lisansı müzakeresi tamamlanmadan kesinleşmez.")
caption("Tablo 7  Önerilen iş paketleri")
table(["Paket ve dönem", "Teknik teslimat", "Önerilen sorumluluk"], [
    ["İP1  Ay 1–3", "BIST ürün/veri envanteri, haklar, veri sözleşmesi, ön kayıt", "Üniversite yöntem; potansiyel paydaş veri/ürün; geliştirici adaptör"],
    ["İP2  Ay 3–6", "Nedensel sinyal ve bar/tick veri kalite harness’i", "Geliştirici altyapı; üniversite bağımsız test"],
    ["İP3  Ay 5–9", "Audit, ücret/gerçekleşme modeli, olay düzeyi MT5/BIST karşılaştırması", "Ortak mühendislik ve piyasa uzmanı doğrulaması"],
    ["İP4  Ay 7–12", "Dondurulmuş strateji ve rejim kampanyası; ML baz çizgileri", "Üniversite deney tasarımı; geliştirici deney koşumu"],
    ["İP5  Ay 10–15", "Kanal/risk fault-injection ve izin varsa test ortamı prototipi", "Geliştirici kontrol katmanı; üye/kurum teknik onayı"],
    ["İP6  Ay 14–18", "İleri dönem gölge izleme, negatif bulgular, makale ve ürün yol haritası", "Üniversite yayın; tüm taraflar hak ve sonuç gözden geçirme"],
], [1.05, 3.25, 2.3])
heading("10.1 Üniversitenin özgün katkısı", 2)
add("Üniversite araştırma ekibi veri sızıntısı ve eşzamanlı bağımlılık için deneysel yöntem, çoklu test düzeltmesi, belirsizlik aralıkları ve strateji seçimi protokolünü tanımlar. Kodun bulunması akademik katkı için yeterli değildir; yayınlanabilir sonuç, tekrar üretilebilir benchmark ve mümkünse anonimleştirilmiş bir hata/olay veri kümesidir.")
heading("10.2 Borsa İstanbul için olası değer", 2)
add("Borsa İstanbul’un olası katkısı, uygun izinlerle ürün tanımları, veri kalitesi açıklamaları, test ortamı bilgisi ve piyasa işleyişi uzmanlığıdır. Olası çıktı; şeffaf backtest kalitesi ölçütleri, eğitim materyali ve algoritmik işlem emniyeti için test senaryolarıdır. Kurumun veri sağlama, finansman, ortaklık veya sonuçları onaylama yükümlülüğü bu metinle oluşmaz.")

heading("11 Veri hakkı, fikri mülkiyet ve yayına hazırlık")
add("Mevcut EA kodları, strateji kütüphaneleri, model artefaktları ve kullanıcıya ait geçmiş kayıtlar proje öncesi arka plan fikri mülkiyeti olarak envanterlenmelidir. BIST verisi ayrı lisans ve erişim koşullarına tabidir; açık kaynak rapora ham verinin taşınabileceği varsayılmaz. Proje sırasında geliştirilen adaptör, audit iyileştirmesi, deney protokolü ve model ağırlıkları için sahiplik, kullanım ve yayın hakları sözleşmede ayrı maddelerdir.")
add("Akademik yayında kod ve veri sürümü, deneme sayısı, maliyet varsayımı, dışlanan işlemler, başarısız rejimler ve finansal çıkar çatışması açıklanır. Veri yeniden dağıtımı mümkün değilse deterministik sentetik örnekler ve çalıştırılabilir değerlendirme kodu paylaşılır. Strateji parametrelerinin veya canlı emir kontrollerinin tamamını yayımlamak bilimsel tekrar üretilebilirliğin tek yolu değildir; erişim kontrollü denetim paketi önerilir.")

heading("12 Başvuru yolu ve karar kapıları")
add("TÜBİTAK 1505 sayfası başvuruların sürekli açık olduğunu, müşteri kuruluş ile yürütücü kuruluşun birlikte başvurduğunu ve başvuru öncesi iş birliği sözleşmesini açıklar [1]. Bu nedenle ilk resmî adım, uygun üniversite araştırma grubu ile Türkiye’de yerleşik müşteri kuruluş rolünü netleştirmektir. Borsa İstanbul potansiyel paydaş olabilir; müşteri kuruluşun mutlaka Borsa İstanbul olması gerektiği varsayılmamalıdır. 1711 Yapay Zekâ Ekosistem 2026 çağrısında finans teknolojileri öncelikliydi fakat ilan edilen son başvuru 18 Eylül 2026 idi; bu rapor tarihi itibarıyla söz konusu çağrı kapanmıştır [9].")
caption("Tablo 8  Başvuru öncesi karar kapıları")
table(["Kapı", "Gerekli karar", "Karar alınmadan yapılmayacak iddia"], [
    ["K1 Ortaklık", "Üniversite ve müşteri kuruluşun rolü, sözleşme", "Kurumlar proje ortağıdır"],
    ["K2 Veri", "BIST veri paketi, lisans, tarihsel kapsama", "BIST üzerinde 10 yıllık test yapıldı"],
    ["K3 Bilimsel ön kayıt", "Hipotez, aday sayısı, final dönem, maliyet", "Seçilmiş strateji kanıtlanmış alpha"],
    ["K4 Teknik emniyet", "Test ortamı, kullanıcı ve sertifika koşulları", "Canlı BIST emir hattı hazır"],
    ["K5 Bütçe", "Gerçek ücret, personel, donanım ve veri maliyeti", "Kesin destek tutarı / yatırım getirisi"],
], [1.1, 2.65, 2.85])
add("Önerilen ilk görüşme paketi bu rapor, kod kanıt dizini, kısa canlı olmayan demo, tek ürün için veri talep listesi ve üç ön kayıtlı deneyden oluşur. Başvuru formunun son dili, ortakların hukuki kimliği ve veri/ürün kapsamı kesinleştikten sonra hazırlanmalıdır.")

heading("Ek A Kaynak kod kanıt dizini")
add("Aşağıdaki [Kxx] kayıtları incelenen çalışma alanındaki dosya ve araştırma notlarını gösterir. Konumlar belge tarihi itibarıyla geçerlidir; canlı terminalde etkin sürüm veya tam fonksiyon doğrulaması oldukları iddia edilmez. Bazı README’ler eski kalmış olabilir; deneyde sabitlenmiş commit ve derleme çıktısı kullanılmalıdır.")
caption("Tablo 9  Teknik iddiaların izlenebilirliği")
table(["Kod", "Çalışma alanı yolu", "Bu raporda desteklediği kapsam"], [
    ["K01", "IndicatorCodeExamples/LiqNarrativeSol_Advanced.mq5", "Liq strateji, RF/rejim, audit, bridge girişleri"],
    ["K02", "IndicatorCodeExamples/OrbUltimateEA_StrategyLibrary.mq5", "ORB seans, NAS filtre, strateji/ML arayüzü"],
    ["K03", "IndicatorCodeExamples/NasUltimate_SATS_PartialNoBE_EA.mq5", "NAS sinyal ve işlem ailesi"],
    ["K04", "IndicatorCodeExamples/MaCrossOverEA.mq5", "MA araştırma ailesi"],
    ["K05", "IndicatorCodeExamples/OrbUltimate_cTrader_Listener.cs", "cTrader listener örneği"],
    ["K06", "IndicatorCodeExamples/NinjaTrader/OrbUltimate_NinjaTrader_Listener.cs", "NinjaTrader listener örneği"],
    ["K07", "IndicatorCodeExamples/Quantower/OrbUltimate_Quantower_Listener.cs", "Quantower listener örneği"],
    ["K08", "IndicatorCodeExamples/macrossover/README.md", "MA audit varsayımı ve kapsam sınırı"],
    ["K09", "docs/MT5_EVENT_VALIDATION_PIPELINE.md", "İç audit ile tester olay karşılaştırma yöntemi"],
    ["K10", "docs/NAS_StrategyLibrary_Implementation.md", "NAS araştırma sınırları ve reddedilen modeller"],
    ["K11", "docs/ORB_ML_Research_And_Validation.md", "Strateji bazlı özellik, ayrım ve yayın kapıları"],
    ["K12", "docs/ORB_Current_vs_R7_Comparison.md", "49 strateji karşılaştırması ve negatif medyanlar"],
    ["K13", "docs/ORB_ChampionCalibration_R8_Results.md", "R8 seçilim sınırı ve çoklu test bulgusu"],
], [0.55, 3.35, 2.7])

heading("Kaynakça")
for item in [
    "[1] TÜBİTAK, 1505 Üniversite-Sanayi İşbirliği Destek Programı. https://tubitak.gov.tr/tr/destekler/sanayi/ulusal-destek-programlari/1505-universite-sanayi-isbirligi-destek-programi",
    "[2] Borsa İstanbul, Veri Yayını. https://www.borsaistanbul.com/veriler/veri-yayini",
    "[3] Borsa İstanbul, Geçmişe Dönük Veri Satışı. https://www.borsaistanbul.com/veriler/gecmise-donuk-veri-satisi",
    "[4] Borsa İstanbul, Test Ortamları. https://www.borsaistanbul.com/teknoloji/teknik-kaynaklar/test-ortamlari",
    "[5] Borsa İstanbul, Bağlantı Protokolleri. https://borsaistanbul.com/bistech-teknolojisi/baglanti-protokolleri",
    "[6] G. C. Cawley ve N. L. C. Talbot, On Over-fitting in Model Selection and Subsequent Selection Bias in Performance Evaluation, JMLR 11, 2010. https://www.jmlr.org/papers/v11/cawley10a.html",
    "[7] D. H. Bailey ve M. López de Prado, The Deflated Sharpe Ratio, 2014. https://doi.org/10.2139/ssrn.2460551",
    "[8] Borsa İstanbul, BISTECH İşlem Öncesi Risk Yönetimi Uygulaması. https://www.borsaistanbul.com/teknoloji/uygulama-yazilimlari",
    "[9] TÜBİTAK, 1711 Yapay Zekâ Ekosistem 2026 Çağrısı. https://tubitak.gov.tr/tr/duyuru/1711-yapay-zeka-ekosistem-2026-yili-cagrisi-acildi",
]:
    add(item)
add("Dış kaynaklara erişim tarihi: 23 Eylül 2026. Program uygunluğu, son çağrı ve mali limitler resmî başvuru öncesinde yeniden doğrulanmalıdır.")

doc.save(OUT)
print(OUT)
